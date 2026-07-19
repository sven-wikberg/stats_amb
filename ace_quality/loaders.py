"""Chargement défensif des anciens CSV et du nouvel export Attrib."""

from __future__ import annotations

import csv
from datetime import date, datetime, time, timedelta
from pathlib import Path
import re
import unicodedata
from typing import Any, Iterable, Iterator

from .models import Intervention


class DataFormatError(ValueError):
    """Signale un fichier absent, incomplet ou d'un format non reconnu."""


def normalize_header(value: Any) -> str:
    """Uniformise accents, espaces et casse sans modifier la donnée originale."""

    text = "" if value is None else str(value)
    text = unicodedata.normalize("NFKD", text)
    text = "".join(char for char in text if not unicodedata.combining(char))
    return re.sub(r"\s+", " ", text).strip().lower()


def _as_text(value: Any) -> str:
    """Convertit proprement les cellules CSV ou Excel en texte exploitable."""

    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value)).strip()


def _parse_date(value: Any) -> date | None:
    """Accepte les dates Excel ainsi que les formats usuels des deux exports."""

    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value

    text = _as_text(value)
    for pattern in ("%d.%m.%Y", "%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def _parse_time(value: Any) -> time | None:
    """Accepte HH:MM, HH:MM:SS, objets Excel et fractions de journée."""

    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    if isinstance(value, (int, float)) and 0 <= value < 1:
        seconds = int(round(float(value) * 24 * 60 * 60)) % (24 * 60 * 60)
        return time(seconds // 3600, (seconds % 3600) // 60, seconds % 60)

    text = _as_text(value)
    for pattern in ("%H:%M:%S", "%H:%M"):
        try:
            return datetime.strptime(text, pattern).time()
        except ValueError:
            continue
    return None


def _parse_duration(value: Any) -> timedelta | None:
    """Convertit une durée Attrib, y compris une durée Excel numérique."""

    if isinstance(value, timedelta):
        return value if value.total_seconds() >= 0 else None
    if isinstance(value, (int, float)) and value >= 0:
        return timedelta(days=float(value))
    if isinstance(value, time):
        return timedelta(hours=value.hour, minutes=value.minute, seconds=value.second)

    text = _as_text(value)
    match = re.fullmatch(r"(?:(\d+)\s+days?,?\s*)?(\d{1,3}):(\d{2})(?::(\d{2}))?", text)
    if not match:
        return None
    days, hours, minutes, seconds = match.groups()
    return timedelta(
        days=int(days or 0),
        hours=int(hours),
        minutes=int(minutes),
        seconds=int(seconds or 0),
    )


def _duration_between(start: time | None, end: time | None) -> timedelta | None:
    """Calcule une durée horaire en gérant explicitement le passage à minuit."""

    if start is None or end is None:
        return None
    reference = date(2000, 1, 1)
    start_dt = datetime.combine(reference, start)
    end_dt = datetime.combine(reference, end)
    if end_dt < start_dt:
        end_dt += timedelta(days=1)
    return end_dt - start_dt


def _parse_int(value: Any, minimum: int = 0, maximum: int | None = None) -> int | None:
    """Lit les nombres entiers sans laisser une valeur aberrante casser le rapport."""

    text = _as_text(value).replace(",", ".")
    try:
        number = int(float(text))
    except (TypeError, ValueError):
        return None
    if number < minimum or (maximum is not None and number > maximum):
        return None
    return number


def _is_yes(value: Any) -> bool:
    """Reconnaît les variantes françaises courantes de Oui/Vrai."""

    return normalize_header(value) in {"oui", "o", "true", "vrai", "1"}


def _clean_vehicle(value: Any) -> str:
    """Affiche ACE708 ou 60708 sous la forme courte 708."""

    text = _as_text(value)
    match = re.search(r"(\d{3})$", text)
    return match.group(1) if match else text


def _raw_priority(row: dict[str, Any]) -> str:
    """Restaure la distinction historique des S1 avec/sans signaux prioritaires."""

    priority = _as_text(row.get("priorite a l'engagement"))
    if priority == "S1":
        signals = _is_yes(row.get("signaux prioritaires jusqu'au site"))
        return "S1 feux bleus" if signals else "S1 sans feux bleus"
    return priority


def _legacy_base(schedule: str) -> tuple[str, str]:
    """Traduit les anciens codes J1/N1/P3 vers une base et une période."""

    schedule = _as_text(schedule)
    mapping = {
        "J1": ("Perréard", "Jour"),
        "N1": ("Perréard", "Nuit"),
        "J2": ("Pierre-du-Niton", "Jour"),
        "N2": ("Pierre-du-Niton", "Nuit"),
        "J3": ("Grangettes Urgences", "Jour"),
        "N3": ("Grangettes Urgences", "Nuit"),
        "P3": ("Grangettes P3", "Jour"),
        "NP3": ("Grangettes P3", "Nuit"),
        "En base CB": ("Perréard", ""),
        "En base PDN": ("Pierre-du-Niton", ""),
        "En base Grangettes": ("Grangettes", ""),
        "NEn base Grangettes": ("Grangettes", "Nuit"),
    }
    return mapping.get(schedule, (schedule, ""))


def _rows_from_csv(path: Path, delimiter: str) -> tuple[list[Any], Iterator[list[Any]]]:
    """Ouvre un CSV UTF-8, avec un repli CP1252 pour d'anciens exports Excel."""

    for encoding in ("utf-8-sig", "cp1252"):
        try:
            # Le fichier est matérialisé ici afin de fermer immédiatement le
            # descripteur, y compris si le schéma est rejeté avant lecture des lignes.
            with path.open(encoding=encoding, newline="") as handle:
                parsed_rows = list(csv.reader(handle, delimiter=delimiter))
            if not parsed_rows:
                raise DataFormatError(f"Le fichier est vide : {path}")
            return parsed_rows[0], iter(parsed_rows[1:])
        except UnicodeDecodeError:
            continue
    raise DataFormatError(f"Encodage non reconnu pour le CSV : {path}")


def _rows_from_xlsx(path: Path) -> tuple[list[Any], Iterator[list[Any]]]:
    """Lit un export XLSX en mode lecture seule pour limiter la mémoire utilisée."""

    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise DataFormatError(
            "La lecture XLSX nécessite openpyxl. Installez les dépendances avec "
            "'python -m pip install -r requirements.txt'."
        ) from exc

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook[workbook.sheetnames[0]]
    rows = worksheet.iter_rows(values_only=True)
    try:
        headers = list(next(rows))
    except StopIteration as exc:
        workbook.close()
        raise DataFormatError(f"Le classeur est vide : {path}") from exc
    # La matérialisation garantit la fermeture du classeur même lorsque la
    # validation des en-têtes échoue avant que l'itérateur ne soit consommé.
    parsed_rows = [list(row) for row in rows]
    workbook.close()
    return headers, iter(parsed_rows)


def _indexed_rows(headers: list[Any], rows: Iterable[list[Any]]) -> Iterator[dict[str, Any]]:
    """Construit des dictionnaires par en-tête, même si l'export a 174 colonnes."""

    positions = {normalize_header(header): index for index, header in enumerate(headers)}
    for row in rows:
        if not any(_as_text(value) for value in row):
            continue
        yield {
            name: row[index] if index < len(row) else None
            for name, index in positions.items()
        }


def _validate_raw_headers(headers: list[Any]) -> None:
    """Échoue tôt avec la liste exacte des colonnes manquantes."""

    available = {normalize_header(header) for header in headers}
    required = {
        "date",
        "priorite a l'engagement",
        "periode du jour",
        "base",
        "vehicule",
        "leader",
        "equipier",
        "heure alarme",
        "depart",
        "arrivee sur site",
        "depart du site",
        "arrivee a destination",
        "temps sur site",
        "probleme principal",
        "naca",
        "patient - age",
    }
    missing = sorted(required - available)
    if missing:
        raise DataFormatError(
            "Le nouvel export ne contient pas les colonnes requises : " + ", ".join(missing)
        )


def _load_raw(headers: list[Any], rows: Iterable[list[Any]]) -> list[Intervention]:
    """Convertit le format Attrib 2026 en interventions normalisées."""

    _validate_raw_headers(headers)
    output: list[Intervention] = []
    for source_row, row in enumerate(_indexed_rows(headers, rows), start=2):
        on_site = _parse_time(row.get("arrivee sur site"))
        depart_site = _parse_time(row.get("depart du site"))
        duration = _parse_duration(row.get("temps sur site"))
        if duration is None:
            duration = _duration_between(on_site, depart_site)

        output.append(
            Intervention(
                source_row=source_row,
                date=_parse_date(row.get("date")),
                status=_as_text(row.get("statut")),
                has_inconsistency=_is_yes(row.get("incoherences")),
                period_of_day=_as_text(row.get("periode du jour")),
                team=_as_text(row.get("equipe")),
                base=_as_text(row.get("base")),
                priority=_raw_priority(row),
                leader=_as_text(row.get("leader")),
                teammate=_as_text(row.get("equipier")),
                vehicle=_clean_vehicle(row.get("vehicule")),
                alarm_at=_parse_time(row.get("heure alarme")),
                departure_at=_parse_time(row.get("depart")),
                on_site_at=on_site,
                depart_site_at=depart_site,
                arrival_destination_at=_parse_time(row.get("arrivee a destination")),
                operational_at=_parse_time(row.get("operationnel")),
                on_site_duration=duration,
                problem=_as_text(row.get("probleme principal")),
                naca=_parse_int(row.get("naca"), minimum=0, maximum=9),
                age=_parse_int(row.get("patient - age"), minimum=0, maximum=120),
                pickup_type=_as_text(row.get("prise en charge - type d'adresse")),
                pickup_place=_as_text(row.get("prise en charge - lieu")),
                pickup_locality=_as_text(row.get("prise en charge - localite")),
                destination_place=_as_text(row.get("destination - lieu")),
                no_transport=_is_yes(row.get("sans transport")),
            )
        )
    return output


def _load_legacy(headers: list[Any], rows: Iterable[list[Any]]) -> list[Intervention]:
    """Maintient la compatibilité avec les fichiers historiques à 33 colonnes."""

    if len(headers) < 33:
        raise DataFormatError(
            f"Ancien format incomplet : 33 colonnes attendues, {len(headers)} trouvées."
        )

    output: list[Intervention] = []
    for source_row, raw_row in enumerate(rows, start=2):
        row = list(raw_row) + [""] * max(0, 33 - len(raw_row))
        if not any(_as_text(value) for value in row):
            continue
        base, inferred_period = _legacy_base(_as_text(row[5]))
        on_site = _parse_time(row[15])
        depart_site = _parse_time(row[16])
        output.append(
            Intervention(
                source_row=source_row,
                date=_parse_date(row[0]),
                period_of_day=_as_text(row[2]) or inferred_period,
                team=_as_text(row[5]),
                base=base,
                priority=_as_text(row[6]),
                leader=_as_text(row[7]),
                teammate=_as_text(row[8]),
                vehicle=_clean_vehicle(row[10]),
                alarm_at=_parse_time(row[13]),
                departure_at=_parse_time(row[14]),
                on_site_at=on_site,
                depart_site_at=depart_site,
                arrival_destination_at=_parse_time(row[17]),
                operational_at=_parse_time(row[18]),
                on_site_duration=_duration_between(on_site, depart_site),
                problem=_as_text(row[24]),
                naca=_parse_int(row[26], minimum=0, maximum=9),
                age=_parse_int(row[32], minimum=0, maximum=120),
                pickup_type=_as_text(row[19]),
                pickup_locality=_as_text(row[20]),
                destination_place=_as_text(row[21]),
                no_transport=not bool(_as_text(row[17])),
            )
        )
    return output


def load_interventions(path: str | Path, delimiter: str = ";") -> list[Intervention]:
    """Détecte automatiquement CSV/XLSX puis ancien/nouveau schéma."""

    source = Path(path)
    if not source.exists():
        raise DataFormatError(f"Fichier de données introuvable : {source}")

    suffix = source.suffix.lower()
    if suffix == ".csv":
        headers, rows = _rows_from_csv(source, delimiter)
    elif suffix == ".xlsx":
        headers, rows = _rows_from_xlsx(source)
    else:
        raise DataFormatError(
            f"Extension non prise en charge '{source.suffix}'. Utilisez .csv ou .xlsx."
        )

    normalized_headers = {normalize_header(header) for header in headers}
    if "priorite a l'engagement" in normalized_headers:
        interventions = _load_raw(headers, rows)
    elif len(headers) >= 33:
        interventions = _load_legacy(headers, rows)
    else:
        raise DataFormatError(
            "Format de données non reconnu : ni export Attrib, ni ancien CSV 33 colonnes."
        )

    if not interventions:
        raise DataFormatError(f"Aucune intervention trouvée dans : {source}")
    return interventions
